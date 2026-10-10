import openpyxl

class EXCEL:

    # def rd_excel(self, file_path):
    #     """读取excel文件"""
    #     try:
    #         workbook = openpyxl.load_workbook(file_path)
    #         return workbook
    #     except Exception as err:
    #         print('读取文件错误：{}'.format(err))

    def rd_excel(self, file_path, sheet_name):
        """
        读取文件
        """
        try:
            workbook = openpyxl.load_workbook(file_path)
            # file_path = '..\\data\\StarRail_data.xlsx'
            # EXCEL = excel_action.EXCEL()
            # workbook = EXCEL.rd_excel(file_path)
            sheet = workbook[sheet_name]
            data1 = []
            data2 = []
            for clo in range(sheet.max_column):
                data1.append(sheet.cell(row=1, column=clo + 1).value)
            for row in range(1, sheet.max_row):
                data3 = {}
                for col in range(sheet.max_column):
                    data3[data1[col]] = sheet.cell(row=row + 1, column=col + 1).value
                    # data3[data1[col]] = sheet.cell(row=row+1,column=col+1).value
                data2.append(data3)
            return data2
            # return workbook
        except Exception as err:
            print('读取文件错误：{}'.format(err))
